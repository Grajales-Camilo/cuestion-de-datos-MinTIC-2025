"""Tipologias territoriales del DNP: parseo y upsert idempotente (T-207,
data-model.md #6, research.md #18).

Fuente: archivo "Base de Datos y Resultados" que el DNP publica cada
vigencia en colaboracion.dnp.gov.co (sin URL estable ano a ano). Columnas
confirmadas contra el archivo real de la vigencia 2026
(`01_ResultadosTipologias2026.xlsx`): hoja `Municipios` usa `CodDANE_txt`
(codigo DIVIPOLA de 5 digitos) como llave; hoja `Departamentos` usa
`cod_dep_txt_corto` (2 digitos) -- NO `cod_dep_txt`, que trae un codigo de
5 digitos distinto al de `divipola_entries.code`. Ambas hojas comparten los
nombres `Tipología_2026`, `Cat 617_2025`, `Poblacion_2024`,
`IngresosTot_2024`. La fila de encabezados no esta en la misma posicion en
ambas hojas (fila 2 en Municipios, fila 3 en Departamentos): se ubica
buscando la columna clave, no por indice fijo.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import select, text
from sqlalchemy.dialects.postgresql import insert as pg_insert
from sqlalchemy.ext.asyncio import AsyncEngine, async_sessionmaker

from app.db.models import DivipolaEntry, TerritorioTipologia

_TIPOLOGIA_COL = "Tipología_2026"
_CATEGORIA_COL = "Cat 617_2025"
_POBLACION_COL = "Poblacion_2024"
_INGRESOS_COL = "IngresosTot_2024"

MUNICIPIO_KEY_COL = "CodDANE_txt"
DEPARTAMENTO_KEY_COL = "cod_dep_txt_corto"

_HEADER_SCAN_ROWS = 5


class TipologiaRowError(ValueError):
    """Fila fuente sin codigo DIVIPOLA o sin tipologia (columna clave vacia)."""


class TipologiaSheetError(ValueError):
    """La hoja no trae la columna clave esperada en las primeras filas."""


@dataclass
class TerritorioTipologiaData:
    divipola_code: str
    level: str
    tipologia_dnp: str
    categoria_ley_617: str | None
    poblacion: int | None
    ingresos_totales_cop: float | None
    vigencia: int
    fuente_archivo: str


def _coerce_number(value: object) -> float | None:
    """Errores de Excel (`#N/A`) y celdas vacias -> None; nunca lanza."""

    if value is None:
        return None
    if isinstance(value, str) and not value.strip():
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _row_to_entry(
    row: dict,
    *,
    key_column: str,
    level: str,
    vigencia: int,
    fuente_archivo: str,
) -> TerritorioTipologiaData:
    divipola_code = str(row.get(key_column) or "").strip()
    tipologia = str(row.get(_TIPOLOGIA_COL) or "").strip()
    if not divipola_code or not tipologia:
        raise TipologiaRowError(f"fila sin codigo DIVIPOLA o tipologia: {row!r}")

    categoria_raw = row.get(_CATEGORIA_COL)
    categoria = str(categoria_raw).strip() if categoria_raw not in (None, "") else None

    poblacion_raw = _coerce_number(row.get(_POBLACION_COL))
    ingresos_raw = _coerce_number(row.get(_INGRESOS_COL))

    return TerritorioTipologiaData(
        divipola_code=divipola_code,
        level=level,
        tipologia_dnp=tipologia,
        categoria_ley_617=categoria,
        poblacion=int(poblacion_raw) if poblacion_raw is not None else None,
        ingresos_totales_cop=ingresos_raw,
        vigencia=vigencia,
        fuente_archivo=fuente_archivo,
    )


def municipio_row_to_entry(
    row: dict, *, vigencia: int, fuente_archivo: str
) -> TerritorioTipologiaData:
    return _row_to_entry(
        row,
        key_column=MUNICIPIO_KEY_COL,
        level="municipality",
        vigencia=vigencia,
        fuente_archivo=fuente_archivo,
    )


def departamento_row_to_entry(
    row: dict, *, vigencia: int, fuente_archivo: str
) -> TerritorioTipologiaData:
    return _row_to_entry(
        row,
        key_column=DEPARTAMENTO_KEY_COL,
        level="department",
        vigencia=vigencia,
        fuente_archivo=fuente_archivo,
    )


@dataclass
class RowsToEntriesResult:
    entries: list[TerritorioTipologiaData]
    skipped_rows: int = 0


def rows_to_entries(
    municipio_rows: list[dict],
    departamento_rows: list[dict],
    *,
    vigencia: int,
    fuente_archivo: str,
) -> RowsToEntriesResult:
    """Convierte filas crudas (dict columna->valor) de ambas hojas en
    entradas `territorio_tipologia`, saltando filas sin codigo/tipologia
    (p. ej. filas en blanco al final de la hoja) sin romper el resto."""

    entries: list[TerritorioTipologiaData] = []
    skipped = 0
    for row in municipio_rows:
        try:
            entries.append(
                municipio_row_to_entry(row, vigencia=vigencia, fuente_archivo=fuente_archivo)
            )
        except TipologiaRowError:
            skipped += 1
    for row in departamento_rows:
        try:
            entries.append(
                departamento_row_to_entry(row, vigencia=vigencia, fuente_archivo=fuente_archivo)
            )
        except TipologiaRowError:
            skipped += 1
    return RowsToEntriesResult(entries=entries, skipped_rows=skipped)


def _find_header(rows: list[tuple], key_column: str) -> tuple[int, dict[str, int]]:
    """Busca, entre las primeras `_HEADER_SCAN_ROWS` filas, la que contiene
    `key_column`, y devuelve `(indice_de_fila, {columna: indice})`."""

    for row_idx, row in enumerate(rows[:_HEADER_SCAN_ROWS]):
        header_index = {
            str(cell).strip(): idx for idx, cell in enumerate(row) if cell is not None
        }
        if key_column in header_index:
            return row_idx, header_index
    raise TipologiaSheetError(
        f"no se encontro la columna clave '{key_column}' en las primeras "
        f"{_HEADER_SCAN_ROWS} filas de la hoja"
    )


def _sheet_rows_to_dicts(rows: list[tuple], key_column: str) -> list[dict]:
    """Ubica el encabezado por nombre de columna (no por indice fijo, ver
    docstring del modulo) y convierte las filas de datos en dicts."""

    header_row_idx, header_index = _find_header(rows, key_column)
    dict_rows = []
    for row in rows[header_row_idx + 1 :]:
        dict_rows.append(
            {column: (row[idx] if idx < len(row) else None) for column, idx in header_index.items()}
        )
    return dict_rows


def read_workbook_entries(path: str | Path, vigencia: int) -> list[TerritorioTipologiaData]:
    """Lee las hojas `Municipios`/`Departamentos` del archivo DNP descargado
    a mano (T-207) y devuelve las entradas listas para
    `upsert_territorio_tipologia`."""

    fuente_archivo = Path(path).name
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        municipio_rows = _sheet_rows_to_dicts(
            list(workbook["Municipios"].iter_rows(values_only=True)), MUNICIPIO_KEY_COL
        )
        departamento_rows = _sheet_rows_to_dicts(
            list(workbook["Departamentos"].iter_rows(values_only=True)), DEPARTAMENTO_KEY_COL
        )
    finally:
        workbook.close()

    result = rows_to_entries(
        municipio_rows,
        departamento_rows,
        vigencia=vigencia,
        fuente_archivo=fuente_archivo,
    )
    return result.entries


@dataclass
class TerritorioTipologiaLoadSummary:
    entries_created: int = 0
    entries_updated: int = 0
    skipped_missing_divipola: list[str] = field(default_factory=list)


async def upsert_territorio_tipologia(
    engine: AsyncEngine, entries: list[TerritorioTipologiaData]
) -> TerritorioTipologiaLoadSummary:
    """Upsert idempotente por `divipola_code`. Entradas cuyo codigo no existe
    (todavia) en `divipola_entries` se reportan en `skipped_missing_divipola`
    y se saltan sin romper la carga completa (T-207)."""

    session_factory = async_sessionmaker(engine, expire_on_commit=False)
    summary = TerritorioTipologiaLoadSummary()
    codes = [entry.divipola_code for entry in entries]

    async with session_factory() as session, session.begin():
        known_divipola_codes = set(
            (
                await session.execute(
                    select(DivipolaEntry.code).where(DivipolaEntry.code.in_(codes))
                )
            ).scalars()
        )

        valid_entries = []
        for entry in entries:
            if entry.divipola_code in known_divipola_codes:
                valid_entries.append(entry)
            else:
                summary.skipped_missing_divipola.append(entry.divipola_code)

        existing_codes = set(
            (
                await session.execute(
                    select(TerritorioTipologia.divipola_code).where(
                        TerritorioTipologia.divipola_code.in_(
                            [entry.divipola_code for entry in valid_entries]
                        )
                    )
                )
            ).scalars()
        )

        for entry in valid_entries:
            stmt = pg_insert(TerritorioTipologia).values(
                divipola_code=entry.divipola_code,
                level=entry.level,
                tipologia_dnp=entry.tipologia_dnp,
                categoria_ley_617=entry.categoria_ley_617,
                poblacion=entry.poblacion,
                ingresos_totales_cop=entry.ingresos_totales_cop,
                vigencia=entry.vigencia,
                fuente_archivo=entry.fuente_archivo,
            )
            stmt = stmt.on_conflict_do_update(
                index_elements=[TerritorioTipologia.divipola_code],
                set_={
                    "level": stmt.excluded.level,
                    "tipologia_dnp": stmt.excluded.tipologia_dnp,
                    "categoria_ley_617": stmt.excluded.categoria_ley_617,
                    "poblacion": stmt.excluded.poblacion,
                    "ingresos_totales_cop": stmt.excluded.ingresos_totales_cop,
                    "vigencia": stmt.excluded.vigencia,
                    "fuente_archivo": stmt.excluded.fuente_archivo,
                    "cargado_at": text("now()"),
                },
            )
            await session.execute(stmt)
            if entry.divipola_code in existing_codes:
                summary.entries_updated += 1
            else:
                summary.entries_created += 1

    return summary
