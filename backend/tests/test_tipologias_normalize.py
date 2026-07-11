import openpyxl
import pytest

from app.db.tipologias import (
    TipologiaRowError,
    TipologiaSheetError,
    departamento_row_to_entry,
    municipio_row_to_entry,
    read_workbook_entries,
    rows_to_entries,
)

_MEDELLIN_ROW = {
    "CodDANE_txt": "05001",
    "Departamento": "ANTIOQUIA",
    "Municipio": "MEDELLÍN",
    "Tipología_2026": "Ciudades grandes",
    "Cat 617_2025": "ESP",
    "Poblacion_2024": 2524747,
    "IngresosTot_2024": 9550875.434461,
}

_CARMEN_VIBORAL_ROW = {
    "CodDANE_txt": "05148",
    "Departamento": "ANTIOQUIA",
    "Municipio": "EL CARMEN DE VIBORAL",
    "Tipología_2026": 5,
    "Cat 617_2025": 6,
    "Poblacion_2024": 18234,
    "IngresosTot_2024": 45123.7,
}

_ANTIOQUIA_ROW = {
    "cod_dep_txt": "05000",
    "cod_dep_txt_corto": "05",
    "Departamento": "Antioquia",
    "Tipología_2026": 1,
    "Cat 617_2025": "ESP",
    "Poblacion_2024": 6903721,
    "IngresosTot_2024": 6362170.081109,
}


def test_municipio_row_to_entry_parses_fields() -> None:
    entry = municipio_row_to_entry(_MEDELLIN_ROW, vigencia=2026, fuente_archivo="x.xlsx")

    assert entry.divipola_code == "05001"
    assert entry.level == "municipality"
    assert entry.tipologia_dnp == "Ciudades grandes"
    assert entry.categoria_ley_617 == "ESP"
    assert entry.poblacion == 2524747
    assert entry.ingresos_totales_cop == pytest.approx(9550875.434461)
    assert entry.vigencia == 2026
    assert entry.fuente_archivo == "x.xlsx"


def test_municipio_row_to_entry_coerces_numeric_tipologia_and_categoria_to_text() -> None:
    entry = municipio_row_to_entry(_CARMEN_VIBORAL_ROW, vigencia=2026, fuente_archivo="x.xlsx")

    assert entry.tipologia_dnp == "5"
    assert entry.categoria_ley_617 == "6"


def test_departamento_row_to_entry_uses_short_code_not_5_digit_code() -> None:
    entry = departamento_row_to_entry(_ANTIOQUIA_ROW, vigencia=2026, fuente_archivo="x.xlsx")

    assert entry.divipola_code == "05"
    assert entry.level == "department"
    assert entry.tipologia_dnp == "1"


@pytest.mark.parametrize(
    "row",
    [
        {"CodDANE_txt": "", "Tipología_2026": "1"},
        {"CodDANE_txt": "05001", "Tipología_2026": ""},
        {"CodDANE_txt": "05001", "Tipología_2026": None},
    ],
)
def test_municipio_row_to_entry_rejects_missing_code_or_tipologia(row: dict) -> None:
    with pytest.raises(TipologiaRowError):
        municipio_row_to_entry(row, vigencia=2026, fuente_archivo="x.xlsx")


def test_municipio_row_to_entry_treats_excel_error_string_as_missing_number() -> None:
    row = {**_MEDELLIN_ROW, "Poblacion_2024": "#N/A", "IngresosTot_2024": None}

    entry = municipio_row_to_entry(row, vigencia=2026, fuente_archivo="x.xlsx")

    assert entry.poblacion is None
    assert entry.ingresos_totales_cop is None


def test_municipio_row_to_entry_missing_categoria_is_none() -> None:
    row = {**_MEDELLIN_ROW, "Cat 617_2025": None}

    entry = municipio_row_to_entry(row, vigencia=2026, fuente_archivo="x.xlsx")

    assert entry.categoria_ley_617 is None


def test_rows_to_entries_combines_and_skips_blank_rows_without_raising() -> None:
    result = rows_to_entries(
        [_MEDELLIN_ROW, _CARMEN_VIBORAL_ROW, {"CodDANE_txt": None, "Tipología_2026": None}],
        [_ANTIOQUIA_ROW],
        vigencia=2026,
        fuente_archivo="x.xlsx",
    )

    assert len(result.entries) == 3
    assert result.skipped_rows == 1
    codes = {entry.divipola_code for entry in result.entries}
    assert codes == {"05001", "05148", "05"}


def _write_workbook(path, *, municipios_header_row: int, departamentos_header_row: int):
    workbook = openpyxl.Workbook()
    municipios = workbook.active
    municipios.title = "Municipios"
    for _ in range(municipios_header_row - 1):
        municipios.append([])
    municipios.append(
        [
            "CodDANE_txt",
            "Municipio",
            "Tipología_2026",
            "Cat 617_2025",
            "Poblacion_2024",
            "IngresosTot_2024",
        ]
    )
    municipios.append(["05001", "MEDELLÍN", "Ciudades grandes", "ESP", 2524747, 9550875.4])
    municipios.append(["05148", "EL CARMEN DE VIBORAL", 5, 6, 18234, 45123.7])

    departamentos = workbook.create_sheet("Departamentos")
    for _ in range(departamentos_header_row - 1):
        departamentos.append([])
    departamentos.append(
        [
            "cod_dep_txt",
            "cod_dep_txt_corto",
            "Departamento",
            "Tipología_2026",
            "Cat 617_2025",
            "Poblacion_2024",
            "IngresosTot_2024",
        ]
    )
    departamentos.append(["05000", "05", "Antioquia", 1, "ESP", 6903721, 6362170.1])

    workbook.save(path)


def test_read_workbook_entries_finds_header_at_different_row_per_sheet(tmp_path) -> None:
    """Reproduce el archivo real del DNP: el encabezado de `Municipios` esta
    en la fila 1 y el de `Departamentos` en la fila 2 (no la misma
    posicion) -- la lectura debe ubicarlo por nombre de columna, no por
    indice fijo."""

    path = tmp_path / "tipologias.xlsx"
    _write_workbook(path, municipios_header_row=1, departamentos_header_row=2)

    entries = read_workbook_entries(path, vigencia=2026)

    codes = {entry.divipola_code for entry in entries}
    assert codes == {"05001", "05148", "05"}
    assert all(entry.fuente_archivo == "tipologias.xlsx" for entry in entries)
    medellin = next(entry for entry in entries if entry.divipola_code == "05001")
    assert medellin.level == "municipality"
    assert medellin.tipologia_dnp == "Ciudades grandes"
    antioquia = next(entry for entry in entries if entry.divipola_code == "05")
    assert antioquia.level == "department"


def test_read_workbook_entries_raises_clear_error_when_key_column_missing(tmp_path) -> None:
    workbook = openpyxl.Workbook()
    municipios = workbook.active
    municipios.title = "Municipios"
    municipios.append(["columna_inesperada"])
    workbook.create_sheet("Departamentos").append(["cod_dep_txt_corto", "Tipología_2026"])
    path = tmp_path / "sin_columna_clave.xlsx"
    workbook.save(path)

    with pytest.raises(TipologiaSheetError):
        read_workbook_entries(path, vigencia=2026)
